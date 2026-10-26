"""Export: a Markdown feedback report and a DOCX copy of the revised resume."""

from __future__ import annotations

import io

from .schemas import Analysis, MatchStatus
from .scoring import score_band

STATUS_LABEL = {
    MatchStatus.MATCHED: "Shown",
    MatchStatus.PARTIAL: "Partly shown",
    MatchStatus.MISSING: "Not shown yet",
}


def build_markdown_report(analysis: Analysis, target_role: str = "") -> str:
    role = target_role or analysis.job.role_title or "this job"
    lines = [
        f"# Resume check: {role}",
        "",
        f"**Fit estimate: {analysis.score}% ({score_band(analysis.score)})**  ",
        f"Keyword coverage: {round(analysis.keyword_coverage * 100)}%  ",
        f"_Scoring {analysis.scoring_version}. An estimate of how well your resume shows the job's "
        "requirements. It is not a hiring chance._",
        "",
        "## Do these first",
    ]
    lines += [f"{i}. **{a.title}** ({a.priority.value}). {a.detail}"
              for i, a in enumerate(analysis.plan.actions, 1)]
    lines += ["", "## Requirement by requirement", "",
              "| # | Requirement | Need | Status | Evidence |", "|---|---|---|---|---|"]
    for req in analysis.job.requirements:
        match = analysis.match_for(req.id)
        status = STATUS_LABEL[match.status] if match else STATUS_LABEL[MatchStatus.MISSING]
        evidence = "; ".join(match.evidence) if match and match.evidence else "None found"
        lines.append(f"| {req.id} | {_cell(req.text)} | {req.priority.value} | {status} | {_cell(evidence)} |")

    if analysis.keywords_missing:
        lines += ["", "## Keywords not in your resume yet", "",
                  ", ".join(analysis.keywords_missing),
                  "", "_Only add a keyword if it's true. Missing here doesn't mean you lack the skill._"]

    if analysis.plan.edits:
        lines += ["", "## Suggested rewrites", ""]
        for edit in analysis.plan.edits:
            lines += [f"- **Before:** {edit.original_text}", f"  **After:** {edit.proposed_text}",
                      f"  _Why:_ {edit.rationale}"]
            lines += [f"  _Check:_ {flag}" for flag in edit.flags]

    if analysis.quality.findings:
        lines += ["", "## Writing notes", ""]
        lines += [f"- **{f.issue}** ({f.severity.value}) {f.location}: {f.why_it_matters}"
                  for f in analysis.quality.findings]

    lines += ["", "## ATS readability", ""]
    lines += [f"- {'Pass' if c.passed else 'Fix'}: **{c.label}**. {c.detail}" for c in analysis.readability]
    lines += ["", "---", "Made with ResumeMania Personal. Placeholders like [N] and [CONFIRM] "
              "mark facts only you can fill in."]
    return "\n".join(lines) + "\n"


def build_resume_docx(resume_text: str) -> bytes:
    """Plain, ATS-friendly DOCX from resume text. Short lines in caps become headings."""
    import docx

    document = docx.Document()
    for raw in resume_text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        if _looks_like_heading(line):
            document.add_heading(line.title() if line.isupper() else line, level=2)
        elif line[:1] in "•-*▪●◦":
            document.add_paragraph(line.lstrip("•-*▪●◦ ").strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _looks_like_heading(line: str) -> bool:
    words = line.split()
    return 0 < len(words) <= 3 and (line.isupper() or line.rstrip(":").lower() in _HEADINGS)


_HEADINGS = {"summary", "skills", "experience", "projects", "education", "certifications",
             "work experience", "technical skills", "achievements"}


def _cell(text: str) -> str:
    return text.replace("|", "/").replace("\n", " ")

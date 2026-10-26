"""Turn an uploaded PDF, DOCX or TXT file into plain text, with input checks."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePath

ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt")
MIN_RESUME_CHARS = 200
MIN_JD_CHARS = 150
MAX_TEXT_CHARS = 40_000


class IngestError(ValueError):
    """User-facing problem with an input. The message is safe to show."""


@dataclass(frozen=True)
class ExtractedText:
    text: str
    source: str  # "pdf" | "docx" | "txt" | "paste"
    page_count: int = 1


def extract_text(filename: str, data: bytes, max_mb: int) -> ExtractedText:
    """Extract text from an uploaded file. Raises IngestError with a clear message."""
    extension = PurePath(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise IngestError("Upload a PDF, DOCX or TXT file.")
    if not data:
        raise IngestError("That file is empty.")
    if len(data) > max_mb * 1024 * 1024:
        raise IngestError(f"That file is over {max_mb} MB. Try a smaller copy.")

    if extension == ".pdf":
        result = _from_pdf(data)
    elif extension == ".docx":
        result = _from_docx(data)
    else:
        result = ExtractedText(text=_decode(data), source="txt")

    return ExtractedText(
        text=clean_resume_text(result.text), source=result.source, page_count=result.page_count
    )


def from_paste(text: str) -> ExtractedText:
    return ExtractedText(text=clean_resume_text(text), source="paste")


def clean_resume_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    collapsed: list[str] = []
    for line in lines:
        if not line and collapsed and not collapsed[-1]:
            continue
        collapsed.append(line)
    return "\n".join(collapsed).strip()[:MAX_TEXT_CHARS]


def validate_inputs(resume_text: str, jd_text: str) -> None:
    """Pre-run hook: refuse to analyse inputs that would give a made-up result."""
    if len(resume_text.strip()) < MIN_RESUME_CHARS:
        raise IngestError(
            "We could only read a little text from your resume. "
            "If it's a scanned PDF, upload a DOCX or paste the text instead."
        )
    if len(jd_text.strip()) < MIN_JD_CHARS:
        raise IngestError("Paste the full job description. A few lines is not enough to match.")


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise IngestError("We couldn't read that text file.")


def _from_pdf(data: bytes) -> ExtractedText:
    import pymupdf

    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise IngestError("That PDF is password protected. Remove the password first.")
            pages = [page.get_text("text") for page in doc]
            return ExtractedText(text="\n".join(pages), source="pdf", page_count=len(pages))
    except IngestError:
        raise
    except Exception as exc:  # pymupdf raises several internal types
        raise IngestError("We couldn't open that PDF. Try saving it again or upload a DOCX.") from exc


def _from_docx(data: bytes) -> ExtractedText:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises zipfile/KeyError variants
        raise IngestError("We couldn't open that DOCX. Try saving it again or upload a PDF.") from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractedText(text="\n".join(parts), source="docx")
